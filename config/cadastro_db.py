from __future__ import annotations

import json
import logging
import os
import sqlite3
from pathlib import Path
from typing import Any

from config.config import DIRETORIOS

logger = logging.getLogger(__name__)

CADASTRO_TEMPLATE = {
    "cod": 0,
    "empresa": "EMPRESA MODELO",
    "cert_path": "caminho/do/certificado.pfx",
    "cert_pass": "senha_do_certificado",
    "cnpj": "cnpj",
    "venc": "",
}

IMPORT_XLSX_ENV = "PORTAL_NFSE_IMPORT_XLSX"
DEFAULT_IMPORT_XLSX = Path.home() / "Desktop" / "Importar.xlsx"


def _db_path() -> Path:
    return Path(DIRETORIOS["cadastros_db"])


def _legacy_json_path() -> Path:
    return Path(DIRETORIOS["cadastros_json"])


def _connect() -> sqlite3.Connection:
    db_path = _db_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def _normalize_cnpj(value: Any) -> str:
    if value is None:
        return ""
    return "".join(ch for ch in str(value).strip() if ch.isalnum())


def _normalize_empresa(data: dict[str, Any]) -> dict[str, Any]:
    cod = int(data["cod"])
    empresa = str(data.get("empresa", "")).strip()
    cnpj = _normalize_cnpj(data.get("cnpj", ""))

    return {
        "cod": cod,
        "empresa": empresa,
        "cnpj": cnpj,
        "cert_path": str(data.get("cert_path", "") or "").strip(),
        "cert_pass": str(data.get("cert_pass", "") or "").strip(),
        "venc": str(data.get("venc", "") or "").strip(),
    }


def _row_to_empresa(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "cod": row["cod"],
        "empresa": row["empresa"],
        "cnpj": row["cnpj"],
        "cert_path": row["cert_path"] or "",
        "cert_pass": row["cert_pass"] or "",
        "venc": row["venc"] or "",
    }


def _is_valid_empresa_record(data: dict[str, Any]) -> bool:
    try:
        empresa = _normalize_empresa(data)
    except (KeyError, TypeError, ValueError):
        return False

    return bool(empresa["empresa"] and empresa["cnpj"])


def initialize_company_database(import_spreadsheet: str | Path | None = None) -> None:
    with _connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS empresas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                cod INTEGER NOT NULL UNIQUE,
                empresa TEXT NOT NULL,
                cnpj TEXT NOT NULL UNIQUE,
                cert_path TEXT NOT NULL DEFAULT '',
                cert_pass TEXT NOT NULL DEFAULT '',
                venc TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE INDEX IF NOT EXISTS idx_empresas_empresa
            ON empresas (empresa);
            """
        )

    if get_company_count() > 0:
        return

    imported = import_companies_from_legacy_json(_legacy_json_path())
    if imported:
        logger.info("Banco de empresas inicializado a partir do cadastro legado em JSON.")
        return

    spreadsheet_path = resolve_import_spreadsheet(import_spreadsheet)
    if spreadsheet_path and spreadsheet_path.exists():
        imported = import_companies_from_spreadsheet(spreadsheet_path, replace_existing=False)
        if imported:
            logger.info("Banco de empresas inicializado a partir da planilha %s.", spreadsheet_path)


def resolve_import_spreadsheet(import_spreadsheet: str | Path | None = None) -> Path | None:
    if import_spreadsheet:
        return Path(import_spreadsheet)

    raw_env_path = os.getenv(IMPORT_XLSX_ENV, "").strip()
    if raw_env_path:
        return Path(raw_env_path)

    return DEFAULT_IMPORT_XLSX


def get_company_count() -> int:
    with _connect() as conn:
        row = conn.execute("SELECT COUNT(*) AS total FROM empresas").fetchone()
        return int(row["total"])


def list_companies() -> list[dict[str, Any]]:
    initialize_company_database()

    with _connect() as conn:
        rows = conn.execute(
            """
            SELECT cod, empresa, cnpj, cert_path, cert_pass, venc
            FROM empresas
            ORDER BY cod
            """
        ).fetchall()

    return [_row_to_empresa(row) for row in rows]


def replace_companies(companies: list[dict[str, Any]]) -> None:
    normalized = []
    seen_codes: set[int] = set()
    seen_cnpjs: set[str] = set()

    for company in companies:
        if not _is_valid_empresa_record(company):
            continue

        normalized_company = _normalize_empresa(company)
        if normalized_company["cod"] in seen_codes or normalized_company["cnpj"] in seen_cnpjs:
            logger.warning(
                "Empresa duplicada ignorada durante sincronizacao com o banco: codigo=%s cnpj=%s",
                normalized_company["cod"],
                normalized_company["cnpj"],
            )
            continue

        seen_codes.add(normalized_company["cod"])
        seen_cnpjs.add(normalized_company["cnpj"])
        normalized.append(normalized_company)

    with _connect() as conn:
        conn.execute("DELETE FROM empresas")
        conn.executemany(
            """
            INSERT INTO empresas (cod, empresa, cnpj, cert_path, cert_pass, venc)
            VALUES (:cod, :empresa, :cnpj, :cert_path, :cert_pass, :venc)
            """,
            normalized,
        )
        conn.commit()


def load_cadastros_payload() -> dict[str, Any]:
    companies = list_companies()
    payload: dict[str, Any] = {
        "cadastros": len(companies) + 1,
        "cadastro_0": CADASTRO_TEMPLATE.copy(),
    }

    for idx, company in enumerate(companies, start=1):
        payload[f"cadastro_{idx}"] = company.copy()

    return payload


def save_cadastros_payload(payload: dict[str, Any]) -> None:
    companies = []

    for key, company in payload.items():
        if not str(key).startswith("cadastro_") or key == "cadastro_0":
            continue
        if isinstance(company, dict):
            companies.append(company)

    replace_companies(companies)


def import_companies_from_legacy_json(json_path: str | Path, replace_existing: bool = False) -> int:
    json_path = Path(json_path)
    if not json_path.exists():
        return 0

    try:
        with json_path.open("r", encoding="utf-8") as file:
            payload = json.load(file)
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Nao foi possivel ler o cadastro legado em JSON: %s", exc)
        return 0

    companies = []
    for key, company in payload.items():
        if not str(key).startswith("cadastro_") or key == "cadastro_0":
            continue
        if isinstance(company, dict) and _is_valid_empresa_record(company):
            companies.append(company)

    if not companies:
        return 0

    if replace_existing or get_company_count() == 0:
        replace_companies(companies)
    else:
        merge_companies(companies)

    return len(companies)


def import_companies_from_spreadsheet(spreadsheet_path: str | Path, replace_existing: bool = False) -> int:
    spreadsheet_path = Path(spreadsheet_path)
    if not spreadsheet_path.exists():
        return 0

    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl nao esta disponivel para importar a planilha de empresas.")
        return 0

    workbook = load_workbook(spreadsheet_path, read_only=True, data_only=True)

    try:
        worksheet = workbook.active
        companies: list[dict[str, Any]] = []

        for row in worksheet.iter_rows(min_col=1, max_col=3, values_only=True):
            raw_cod, raw_empresa, raw_cnpj = row

            if raw_cod in (None, "") and raw_empresa in (None, "") and raw_cnpj in (None, ""):
                continue

            try:
                cod = int(str(raw_cod).strip())
            except (TypeError, ValueError):
                continue

            empresa = str(raw_empresa or "").strip()
            cnpj = _normalize_cnpj(raw_cnpj)

            if not empresa or not cnpj:
                continue

            companies.append(
                {
                    "cod": cod,
                    "empresa": empresa,
                    "cnpj": cnpj,
                    "cert_path": "",
                    "cert_pass": "",
                    "venc": "",
                }
            )
    finally:
        workbook.close()

    if not companies:
        return 0

    if replace_existing or get_company_count() == 0:
        replace_companies(companies)
    else:
        merge_companies(companies)

    return len(companies)


def merge_companies(companies: list[dict[str, Any]]) -> None:
    existing_by_code = {company["cod"]: company for company in list_companies()}

    for company in companies:
        if not _is_valid_empresa_record(company):
            continue

        normalized_company = _normalize_empresa(company)
        current = existing_by_code.get(normalized_company["cod"])

        if current:
            current.update(
                {
                    "empresa": normalized_company["empresa"],
                    "cnpj": normalized_company["cnpj"],
                }
            )
            if normalized_company["cert_path"]:
                current["cert_path"] = normalized_company["cert_path"]
            if normalized_company["cert_pass"]:
                current["cert_pass"] = normalized_company["cert_pass"]
            if normalized_company["venc"]:
                current["venc"] = normalized_company["venc"]
        else:
            existing_by_code[normalized_company["cod"]] = normalized_company

    replace_companies(list(existing_by_code.values()))
