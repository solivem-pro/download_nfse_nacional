from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path

from config.config import DIRETORIOS

logger = logging.getLogger(__name__)

INVALID_FS_CHARS = r'[<>:"/\\|?*]'


def sanitize_company_folder_name(company_name: str | None, cod_empresa: str | int) -> str:
    text = str(company_name or "").strip()
    if text:
        text = re.sub(INVALID_FS_CHARS, " ", text)
        text = " ".join(text.split()).strip(" .")

    return text or f"Empresa {cod_empresa}"


def company_root_dir(cod_empresa: str | int, company_name: str | None) -> Path:
    folder_name = sanitize_company_folder_name(company_name, cod_empresa)
    return Path(DIRETORIOS["notas"]) / folder_name


def company_period_dir(cod_empresa: str | int, company_name: str | None, year: str | int, month: str | int) -> Path:
    return company_root_dir(cod_empresa, company_name) / competence_folder_name(year, month)


def company_control_file(cod_empresa: str | int, company_name: str | None) -> Path:
    return company_root_dir(cod_empresa, company_name) / "nsu_competencia.json"


def competence_folder_name(year: str | int, month: str | int) -> str:
    return f"{int(month):02d}-{int(year):04d}"


def ensure_company_workspace(
    cod_empresa: str | int,
    cnpj: str,
    company_name: str | None,
    previous_company_name: str | None = None,
) -> tuple[bool, str]:
    """
    Garante que a pasta raiz da empresa exista em /packs/<empresa>.

    Mantemos o arquivo de controle NSU nesta pasta raiz, enquanto os artefatos
    de download passam a ser gravados dentro de subpastas por competência.
    """
    cod_str = str(cod_empresa)
    template_dir = Path(DIRETORIOS["notas"]) / "0"
    target_dir = company_root_dir(cod_empresa, company_name)
    legacy_code_dir = Path(DIRETORIOS["notas"]) / cod_str
    previous_name_dir = (
        company_root_dir(cod_empresa, previous_company_name)
        if previous_company_name and previous_company_name != company_name
        else None
    )

    try:
        migration_source = None
        if target_dir.exists():
            migration_source = None
        elif previous_name_dir and previous_name_dir.exists():
            migration_source = previous_name_dir
        elif legacy_code_dir.exists() and legacy_code_dir != target_dir:
            migration_source = legacy_code_dir

        if migration_source is not None and migration_source.exists():
            shutil.move(str(migration_source), str(target_dir))
            logger.info("Pasta da empresa %s migrada de %s para %s.", cod_str, migration_source.name, target_dir.name)

        if not target_dir.exists():
            if template_dir.exists() and any(template_dir.iterdir()):
                shutil.copytree(template_dir, target_dir)
                logger.info("Estrutura da empresa %s criada a partir da pasta modelo.", cod_str)
            else:
                target_dir.mkdir(parents=True, exist_ok=True)
                logger.info("Pasta raiz da empresa %s criada sem modelo base.", cod_str)

        _configure_company_report(target_dir, cnpj, cod_str)
        return True, str(target_dir)

    except Exception as exc:
        logger.error("Erro ao preparar estrutura da empresa %s: %s", cod_str, exc)
        return False, f"Erro ao preparar estrutura da empresa {cod_str}: {exc}"


def ensure_period_workspace(
    cod_empresa: str | int,
    cnpj: str,
    company_name: str | None,
    year: str | int,
    month: str | int,
) -> tuple[bool, str]:
    """
    Garante a pasta da competência:
    /packs/<Empresa>/<MM-AAAA>/{PRESTADOS,TOMADOS,EVENTOS}
    """
    ok, message = ensure_company_workspace(cod_empresa, cnpj, company_name)
    if not ok:
        return False, message

    root_dir = Path(message)
    period_dir = root_dir / competence_folder_name(year, month)

    try:
        period_dir.mkdir(parents=True, exist_ok=True)
        for folder_name in ("PRESTADOS", "TOMADOS", "EVENTOS"):
            folder = period_dir / folder_name
            folder.mkdir(parents=True, exist_ok=True)
            if folder_name in {"PRESTADOS", "TOMADOS"}:
                (folder / "Canceladas").mkdir(parents=True, exist_ok=True)

        _configure_period_report(root_dir, period_dir, cnpj, str(cod_empresa))
        return True, str(period_dir)
    except Exception as exc:
        logger.error("Erro ao preparar pasta da competência %s/%s para a empresa %s: %s", month, year, cod_empresa, exc)
        return False, f"Erro ao preparar pasta da competência {month}-{year} da empresa {cod_empresa}: {exc}"


def _configure_company_report(company_dir: Path, cnpj: str, cod_empresa: str) -> None:
    """Mantém um relatório base na pasta raiz da empresa, quando disponível."""
    expected_report = company_dir / f"relatorio_{cod_empresa}.xlsm"
    report_file = _ensure_report_file(company_dir, expected_report)
    if report_file is None:
        return

    _write_target_cnpj(report_file, cnpj, cod_empresa)


def _configure_period_report(company_dir: Path, period_dir: Path, cnpj: str, cod_empresa: str) -> None:
    expected_report = period_dir / f"relatorio_{cod_empresa}.xlsm"
    report_file = _ensure_report_file(period_dir, expected_report, source_dir=company_dir)
    if report_file is None:
        return

    _write_target_cnpj(report_file, cnpj, cod_empresa)


def _ensure_report_file(target_dir: Path, expected_report: Path, source_dir: Path | None = None) -> Path | None:
    if expected_report.exists():
        return expected_report

    xlsm_files = sorted(target_dir.glob("*.xlsm"))
    if xlsm_files:
        report_file = xlsm_files[0]
        if report_file != expected_report:
            try:
                report_file.rename(expected_report)
                logger.info("Relatório renomeado para %s.", expected_report.name)
                return expected_report
            except OSError as exc:
                logger.warning("Não foi possível renomear o relatório %s: %s", report_file, exc)
                return report_file
        return report_file

    candidate_sources: list[Path] = []
    if source_dir is not None:
        candidate_sources.extend(sorted(source_dir.glob("*.xlsm")))

    model_path = Path(DIRETORIOS["planilha_modelo"])
    if model_path.exists():
        candidate_sources.append(model_path)

    for source in candidate_sources:
        if not source.exists():
            continue
        try:
            shutil.copy2(source, expected_report)
            logger.info("Relatório copiado para %s a partir de %s.", expected_report, source)
            return expected_report
        except OSError as exc:
            logger.warning("Não foi possível copiar o relatório base %s: %s", source, exc)

    logger.warning("Nenhum arquivo .xlsm encontrado para preparar o relatório em %s.", target_dir)
    return None


def _write_target_cnpj(report_file: Path, cnpj: str, cod_empresa: str) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("Biblioteca openpyxl não disponível para ajustar o relatório da empresa %s.", cod_empresa)
        return

    workbook = None
    try:
        workbook = load_workbook(report_file, keep_vba=True)
        if "alvo" not in workbook.sheetnames:
            return

        sheet = workbook["alvo"]
        cnpj_str = str(cnpj)

        if sheet["A1"].value != cnpj_str:
            sheet["A1"] = cnpj_str
            sheet["A1"].number_format = "@"
            workbook.save(report_file)
            logger.info("Relatório da empresa %s atualizado com o CNPJ em alvo!A1.", cod_empresa)
    except Exception as exc:
        logger.warning("Erro ao ajustar o relatório da empresa %s: %s", cod_empresa, exc)
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
