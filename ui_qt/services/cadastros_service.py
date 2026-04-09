from __future__ import annotations

import json
import logging
import shutil
import subprocess
import math
from datetime import datetime
from pathlib import Path
from typing import Any

from config.cadastro_db import load_cadastros_payload, resolve_import_spreadsheet, save_cadastros_payload
from config.company_workspace import company_control_file, company_root_dir, ensure_company_workspace
from config.config import DIRETORIOS, ROOT_DIR
from config.utils import formatar_cnpj, limpar_cnpj, limpar_numero, validar_cnpj

logger = logging.getLogger(__name__)

NSU_RESET_PAYLOAD = {
    "registros": {
        "2000": {
            "01": {"nsu_inicial": 0, "nsu_final": 0},
        }
    }
}


def list_company_forms() -> list[dict[str, Any]]:
    payload = load_cadastros_payload()
    companies = []

    for key, company in payload.items():
        if not str(key).startswith("cadastro_") or key == "cadastro_0" or not isinstance(company, dict):
            continue
        companies.append(_company_from_payload(company))

    companies.sort(key=lambda company: company["cod"])
    return companies


def get_company_by_code(cod: int | str) -> dict[str, Any] | None:
    target = int(cod)
    for company in list_company_forms():
        if int(company["cod"]) == target:
            return company
    return None


def upsert_company(company: dict[str, Any], edit_code: int | None = None) -> dict[str, Any]:
    normalized = _normalize_company_input(company)
    _validate_uniqueness(normalized["cod"], normalized["cnpj"], edit_code=edit_code)

    payload = load_cadastros_payload()
    target_key = None
    previous_name = None

    if edit_code is not None:
        for key, existing in payload.items():
            if not str(key).startswith("cadastro_") or key == "cadastro_0" or not isinstance(existing, dict):
                continue
            if int(existing.get("cod", 0)) == int(edit_code):
                target_key = key
                previous_name = str(existing.get("empresa", "") or "").strip()
                break

    if target_key is None:
        next_index = int(payload.get("cadastros", 1))
        target_key = f"cadastro_{next_index}"
        payload["cadastros"] = next_index + 1
        model = dict(payload.get("cadastro_0", {}))
    else:
        model = dict(payload.get(target_key, {}))

    model.update(normalized)
    payload[target_key] = model
    save_cadastros_payload(payload)

    workspace_ok, workspace_msg = ensure_company_workspace(
        normalized["cod"],
        normalized["cnpj"],
        normalized["empresa"],
        previous_company_name=previous_name,
    )
    if not workspace_ok:
        logger.warning(workspace_msg)

    return _company_from_payload(model)


def delete_company(cod: int | str) -> dict[str, Any] | None:
    target = int(cod)
    payload = load_cadastros_payload()
    removed = None
    remove_key = None

    for key, company in payload.items():
        if not str(key).startswith("cadastro_") or key == "cadastro_0" or not isinstance(company, dict):
            continue
        if int(company.get("cod", 0)) == target:
            removed = _company_from_payload(company)
            remove_key = key
            break

    if remove_key is None:
        return None

    _remove_company_assets(payload[remove_key])
    payload.pop(remove_key, None)
    save_cadastros_payload(payload)
    return removed


def delete_all_companies() -> int:
    companies = list_company_forms()
    for company in companies:
        _remove_company_assets(company)

    payload = load_cadastros_payload()
    payload = {"cadastros": 1, "cadastro_0": dict(payload.get("cadastro_0", {}))}
    save_cadastros_payload(payload)
    return len(companies)


def reset_all_nsu_files() -> int:
    updated = 0
    packs_dir = Path(DIRETORIOS["packs"])
    if not packs_dir.exists():
        return 0

    for path in packs_dir.rglob("nsu_competencia.json"):
        path.write_text(json.dumps(NSU_RESET_PAYLOAD, ensure_ascii=False, indent=2), encoding="utf-8")
        updated += 1

    return updated


def suggested_import_spreadsheet_path() -> Path | None:
    return resolve_import_spreadsheet()


def import_companies_from_xlsx(spreadsheet_path: str | Path) -> dict[str, Any]:
    path = Path(spreadsheet_path)
    if not path.exists():
        raise FileNotFoundError("Planilha .xlsx não encontrada.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("A biblioteca openpyxl não está disponível para importar a planilha.") from exc

    created: list[dict[str, Any]] = []
    updated: list[dict[str, Any]] = []
    unchanged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    existing_by_code = {int(company["cod"]): company for company in list_company_forms()}
    existing_by_cnpj = {
        limpar_cnpj(str(company.get("cnpj", "") or "")): company
        for company in existing_by_code.values()
        if str(company.get("cnpj", "") or "").strip()
    }
    seen_codes: set[int] = set()
    seen_cnpjs: set[str] = set()

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active

        for row_number, row in enumerate(worksheet.iter_rows(min_col=1, max_col=3, values_only=True), start=1):
            parsed = _parse_import_row(row_number, row)
            if parsed is None:
                continue
            if "error" in parsed:
                errors.append({"row": row_number, "message": str(parsed["error"])})
                continue

            cod = int(parsed["cod"])
            empresa = str(parsed["empresa"])
            cnpj = str(parsed["cnpj"])

            if cod in seen_codes:
                errors.append({"row": row_number, "message": f"Código {cod} repetido na planilha."})
                continue
            if cnpj in seen_cnpjs:
                errors.append(
                    {"row": row_number, "message": f"CNPJ {formatar_cnpj(cnpj)} repetido na planilha."}
                )
                continue

            seen_codes.add(cod)
            seen_cnpjs.add(cnpj)

            current = existing_by_code.get(cod)
            cnpj_owner = existing_by_cnpj.get(cnpj)

            if current is None and cnpj_owner is not None:
                errors.append(
                    {
                        "row": row_number,
                        "message": (
                            f"CNPJ {formatar_cnpj(cnpj)} já cadastrado para o código {int(cnpj_owner['cod'])}."
                        ),
                    }
                )
                continue

            if current is not None and cnpj_owner is not None and int(cnpj_owner["cod"]) != int(current["cod"]):
                errors.append(
                    {
                        "row": row_number,
                        "message": (
                            f"CNPJ {formatar_cnpj(cnpj)} já cadastrado para o código {int(cnpj_owner['cod'])}."
                        ),
                    }
                )
                continue

            try:
                if current is None:
                    saved = upsert_company(
                        {
                            "cod": str(cod),
                            "empresa": empresa,
                            "cnpj": cnpj,
                            "cert_pass": "",
                            "cert_path": "",
                            "venc": "",
                        }
                    )
                    created.append(_import_result_item(row_number, saved))
                else:
                    if not _company_changed(current, empresa, cnpj):
                        unchanged.append(_import_result_item(row_number, current))
                        continue

                    saved = upsert_company(
                        {
                            "cod": str(current["cod"]),
                            "empresa": empresa,
                            "cnpj": cnpj,
                            "cert_pass": str(current.get("cert_pass", "") or ""),
                            "cert_path": str(current.get("cert_path", "") or ""),
                            "venc": str(current.get("venc", "") or ""),
                        },
                        edit_code=int(current["cod"]),
                    )
                    updated.append(_import_result_item(row_number, saved))
            except ValueError as exc:
                errors.append({"row": row_number, "message": str(exc)})
                continue
            except Exception as exc:
                errors.append({"row": row_number, "message": f"Falha ao importar a linha: {exc}"})
                continue

            existing_by_code[int(saved["cod"])] = saved
            existing_by_cnpj[limpar_cnpj(str(saved.get("cnpj", "") or ""))] = saved
    finally:
        workbook.close()

    return {
        "path": str(path),
        "sheet": worksheet.title if "worksheet" in locals() else "",
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "errors": errors,
    }


def copy_certificate_for_company(source_file: str | Path, cod: int | str, cert_pass: str) -> tuple[str, str]:
    source = Path(source_file)
    if not source.exists():
        raise FileNotFoundError("Arquivo de certificado nao encontrado.")

    destination = Path(DIRETORIOS["certificados"]) / f"{int(cod)}.pfx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    relative_path = str(destination.relative_to(ROOT_DIR))
    return relative_path, read_certificate_expiry(destination, cert_pass)


def read_certificate_expiry(cert_path: str | Path, cert_pass: str) -> str:
    cert_path = Path(cert_path)
    if not cert_path.exists():
        return ""

    password = cert_pass.encode() if cert_pass else None

    try:
        from cryptography.hazmat.backends import default_backend
        from cryptography.hazmat.primitives.serialization import pkcs12

        cert_bytes = cert_path.read_bytes()
        _, certificate, _ = pkcs12.load_key_and_certificates(cert_bytes, password, default_backend())
        if certificate:
            if hasattr(certificate, "not_valid_after_utc"):
                expiration = certificate.not_valid_after_utc.astimezone().replace(tzinfo=None)
            else:
                expiration = certificate.not_valid_after
            return expiration.strftime("%d/%m/%Y")
    except Exception as exc:
        logger.info("Nao foi possivel ler vencimento do certificado com cryptography: %s", exc)

    try:
        result = subprocess.run(
            [
                "openssl",
                "pkcs12",
                "-in",
                str(cert_path),
                "-clcerts",
                "-nokeys",
                "-passin",
                f"pass:{cert_pass}",
            ],
            capture_output=True,
            text=True,
            timeout=10,
        )
        if result.returncode == 0:
            for line in result.stdout.splitlines():
                if "notAfter=" not in line:
                    continue
                expiry = line.split("=", 1)[1].strip()
                return datetime.strptime(expiry, "%b %d %H:%M:%S %Y GMT").strftime("%d/%m/%Y")
    except Exception as exc:
        logger.info("Nao foi possivel ler vencimento do certificado com openssl: %s", exc)

    return ""


def load_nsu_payload(cod: int | str) -> dict[str, Any]:
    path = _nsu_file_path(cod)
    if not path.exists():
        payload = {"registros": {}}
        save_nsu_payload(cod, payload)
        return payload

    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        logger.warning("Arquivo NSU invalido para a empresa %s; recriando payload vazio.", cod)
        payload = {"registros": {}}
        save_nsu_payload(cod, payload)
        return payload


def save_nsu_payload(cod: int | str, payload: dict[str, Any]) -> None:
    path = _nsu_file_path(cod)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = {"registros": payload.get("registros", {})}
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def nsu_rows_for_company(cod: int | str) -> list[dict[str, Any]]:
    payload = load_nsu_payload(cod)
    rows = []
    for ano, meses in payload.get("registros", {}).items():
        for mes, valores in meses.items():
            rows.append(
                {
                    "ano": str(ano),
                    "mes": str(mes).zfill(2),
                    "nsu_inicial": int(valores.get("nsu_inicial", 0) or 0),
                    "nsu_final": int(valores.get("nsu_final", 0) or 0),
                }
            )
    rows.sort(key=lambda row: (int(row["ano"]), int(row["mes"])), reverse=True)
    return rows


def upsert_nsu_record(cod: int | str, ano: str | int, mes: str | int, nsu_inicial: str | int, nsu_final: str | int) -> None:
    ano_str = str(int(ano))
    mes_str = f"{int(mes):02d}"
    if not 1 <= int(mes) <= 12:
        raise ValueError("M\u00eas deve estar entre 1 e 12.")

    payload = load_nsu_payload(cod)
    payload.setdefault("registros", {})
    payload["registros"].setdefault(ano_str, {})
    payload["registros"][ano_str][mes_str] = {
        "nsu_inicial": int(limpar_numero(nsu_inicial)),
        "nsu_final": int(limpar_numero(nsu_final)),
    }
    save_nsu_payload(cod, payload)


def delete_nsu_record(cod: int | str, ano: str | int, mes: str | int) -> None:
    ano_str = str(int(ano))
    mes_str = f"{int(mes):02d}"
    payload = load_nsu_payload(cod)
    registros = payload.get("registros", {})
    if ano_str in registros and mes_str in registros[ano_str]:
        del registros[ano_str][mes_str]
        if not registros[ano_str]:
            del registros[ano_str]
        save_nsu_payload(cod, payload)


def clear_nsu_records(cod: int | str) -> None:
    save_nsu_payload(cod, {"registros": {}})


def _normalize_company_input(company: dict[str, Any]) -> dict[str, Any]:
    empresa = str(company.get("empresa", "")).strip()
    if not empresa:
        raise ValueError("Informe o nome da empresa.")

    raw_cod = str(company.get("cod", "")).strip()
    if not raw_cod:
        raise ValueError("Informe o c\u00f3digo da empresa.")

    try:
        cod = int(raw_cod)
    except ValueError as exc:
        raise ValueError("C\u00f3digo deve ser um n\u00famero inteiro.") from exc

    cnpj = validar_cnpj(str(company.get("cnpj", "")))
    if not cnpj:
        raise ValueError("CNPJ inv\u00e1lido, deve conter 14 d\u00edgitos alfanum\u00e9ricos.")

    return {
        "cod": cod,
        "empresa": empresa,
        "cnpj": limpar_cnpj(cnpj),
        "cert_pass": str(company.get("cert_pass", "") or "").strip(),
        "cert_path": str(company.get("cert_path", "") or "").strip(),
        "venc": str(company.get("venc", "") or "").strip(),
    }


def _validate_uniqueness(cod: int, cnpj: str, edit_code: int | None = None) -> None:
    for company in list_company_forms():
        company_code = int(company["cod"])
        if edit_code is not None and company_code == int(edit_code):
            continue

        if company_code == int(cod):
            raise ValueError(f"O codigo {cod} ja esta em uso pela empresa {company['empresa']}.")

        if limpar_cnpj(company["cnpj"]) == limpar_cnpj(cnpj):
            raise ValueError(
                f"O CNPJ {formatar_cnpj(cnpj)} ja esta em uso pela empresa {company['empresa']}."
            )


def _company_from_payload(company: dict[str, Any]) -> dict[str, Any]:
    cod = int(company.get("cod", 0))
    cnpj = limpar_cnpj(str(company.get("cnpj", "") or ""))
    return {
        "cod": cod,
        "empresa": str(company.get("empresa", "") or "").strip(),
        "cnpj": cnpj,
        "cnpj_formatado": formatar_cnpj(cnpj),
        "cert_pass": str(company.get("cert_pass", "") or "").strip(),
        "cert_path": str(company.get("cert_path", "") or "").strip(),
        "venc": str(company.get("venc", "") or "").strip(),
    }


def _nsu_file_path(cod: int | str) -> Path:
    company = get_company_by_code(cod)
    company_name = company["empresa"] if company else None
    return company_control_file(cod, company_name)


def _remove_company_assets(company: dict[str, Any]) -> None:
    cert_file = Path(str(company.get("cert_path", "") or ""))
    if cert_file.name:
        absolute_cert = cert_file if cert_file.is_absolute() else Path(ROOT_DIR) / cert_file
        if absolute_cert.exists():
            try:
                absolute_cert.unlink()
            except OSError as exc:
                logger.warning("Nao foi possivel remover o certificado %s: %s", absolute_cert, exc)

    notes_dir = company_root_dir(int(company.get("cod", 0)), str(company.get("empresa", "") or ""))
    if notes_dir.exists():
        try:
            shutil.rmtree(notes_dir)
        except OSError as exc:
            logger.warning("Nao foi possivel remover a pasta %s: %s", notes_dir, exc)

    legacy_notes_dir = Path(DIRETORIOS["notas"]) / str(int(company.get("cod", 0)))
    if legacy_notes_dir.exists() and legacy_notes_dir != notes_dir:
        try:
            shutil.rmtree(legacy_notes_dir)
        except OSError as exc:
            logger.warning("Nao foi possivel remover a pasta legada %s: %s", legacy_notes_dir, exc)

    zip_file = Path(DIRETORIOS["notas"]) / f"{int(company.get('cod', 0))}.zip"
    if zip_file.exists():
        try:
            zip_file.unlink()
        except OSError as exc:
            logger.warning("Nao foi possivel remover o arquivo %s: %s", zip_file, exc)


def _parse_import_row(row_number: int, row: tuple[Any, Any, Any] | tuple[Any, ...]) -> dict[str, Any] | None:
    raw_cod, raw_empresa, raw_cnpj = (list(row) + [None, None, None])[:3]

    if _is_empty_import_row(raw_cod, raw_empresa, raw_cnpj):
        return None

    if row_number == 1 and _looks_like_import_header(raw_cod, raw_empresa, raw_cnpj):
        return None

    cod = _coerce_import_code(raw_cod)
    if cod is None:
        return {"error": "Código inválido ou ausente."}

    empresa = str(raw_empresa or "").strip()
    if not empresa:
        return {"error": "Nome da empresa ausente."}

    cnpj = validar_cnpj(str(raw_cnpj or ""))
    if not cnpj:
        return {"error": "CNPJ inválido ou ausente."}

    return {"cod": cod, "empresa": empresa, "cnpj": cnpj}


def _is_empty_import_row(raw_cod: Any, raw_empresa: Any, raw_cnpj: Any) -> bool:
    values = [raw_cod, raw_empresa, raw_cnpj]
    return all(value in (None, "") or not str(value).strip() for value in values)


def _looks_like_import_header(raw_cod: Any, raw_empresa: Any, raw_cnpj: Any) -> bool:
    tokens = [str(value or "").strip().lower() for value in (raw_cod, raw_empresa, raw_cnpj)]
    text = " ".join(tokens)
    return "cnpj" in text and ("codigo" in text or "código" in text or "empresa" in text)


def _coerce_import_code(raw_cod: Any) -> int | None:
    if raw_cod in (None, ""):
        return None

    if isinstance(raw_cod, int):
        return int(raw_cod)
    if isinstance(raw_cod, float):
        if math.isnan(raw_cod):
            return None
        if raw_cod.is_integer():
            return int(raw_cod)
        digits = "".join(ch for ch in format(raw_cod, "g") if ch.isdigit())
        return int(digits) if digits else None

    text = str(raw_cod).strip()
    if not text:
        return None

    digits = "".join(ch for ch in text if ch.isdigit())
    if digits:
        if any(separator in text for separator in (",", ".")):
            return int(digits)

    try:
        return int(text)
    except ValueError:
        try:
            numeric = float(text.replace(",", "."))
        except ValueError:
            return None
        if not numeric.is_integer():
            return None
        return int(numeric)


def _company_changed(current: dict[str, Any], empresa: str, cnpj: str) -> bool:
    return str(current.get("empresa", "") or "").strip() != empresa or limpar_cnpj(
        str(current.get("cnpj", "") or "")
    ) != limpar_cnpj(cnpj)


def _import_result_item(row_number: int, company: dict[str, Any]) -> dict[str, Any]:
    cnpj = limpar_cnpj(str(company.get("cnpj", "") or ""))
    return {
        "row": row_number,
        "cod": int(company["cod"]),
        "empresa": str(company.get("empresa", "") or "").strip(),
        "cnpj": cnpj,
        "cnpj_formatado": formatar_cnpj(cnpj),
    }
